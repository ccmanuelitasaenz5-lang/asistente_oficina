"""
Exportador de Padrón a Excel
=============================
Genera un archivo .xlsx del padrón completo de voceros con:
  - Hoja resumen con estadísticas por consejo
  - Una hoja por cada Consejo Comunal con tabla formateada
  - Colores institucionales, filtros, bordes profesionales
"""

import re
import logging
from datetime import datetime
from pathlib import Path
from io import BytesIO
from typing import Optional

logger = logging.getLogger("excel_exporter")

# ── Colores institucionales ──────────────────────────────────────────────────
COLOR_HEADER     = "1E3A5F"   # Azul oscuro
COLOR_SUBHEADER  = "2E6DA4"   # Azul medio
COLOR_ALT_ROW    = "EBF3FB"   # Azul muy claro (filas alternas)
COLOR_ACENTO     = "C0392B"   # Rojo para totales
COLOR_BLANCO     = "FFFFFF"
COLOR_BORDE      = "BFBFBF"
FUENTE           = "Arial"


async def _obtener_datos_padron() -> dict:
    """Obtiene todos los voceros desde SQLite, agrupados por consejo comunal."""
    from backend.database import get_db
    consejos = {}
    async with await get_db() as db:
        cursor = await db.execute(
            """
            SELECT
                c.nombre  AS consejo,
                c.sector  AS sector,
                v.nombre  AS nombre,
                v.cedula  AS cedula,
                v.cargo   AS cargo,
                v.tipo    AS tipo,
                v.telefono AS telefono,
                v.votos   AS votos,
                v.email   AS email
            FROM voceros v
            LEFT JOIN consejos_comunales c ON v.consejo_id = c.id
            WHERE v.activo = 1
            ORDER BY c.nombre ASC, v.cargo ASC, v.tipo ASC, v.nombre ASC
            """
        )
        rows = await cursor.fetchall()
        for r in rows:
            nombre_c = (r[0] or "Sin Consejo Asignado").strip()
            if nombre_c not in consejos:
                consejos[nombre_c] = {"sector": (r[1] or "").strip(), "voceros": []}
            consejos[nombre_c]["voceros"].append({
                "nombre":   (r[2] or "").strip(),
                "cedula":   (r[3] or "").strip(),
                "cargo":    (r[4] or "").strip(),
                "tipo":     (r[5] or "Principal").strip(),
                "telefono": (r[6] or "").strip(),
                "votos":    int(r[7] or 0),
                "email":    (r[8] or "").strip(),
            })
    return consejos


def _celda(ws, fila, col, valor,
           negrita=False, color_fondo=None, color_fuente="000000",
           tamaño=10, alineacion="left", borde=False):
    """Aplica valor y formato a una celda."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    cell = ws.cell(row=fila, column=col, value=valor)
    cell.font = Font(name=FUENTE, bold=negrita, size=tamaño, color=color_fuente)
    if color_fondo:
        cell.fill = PatternFill("solid", fgColor=color_fondo)
    cell.alignment = Alignment(horizontal=alineacion, vertical="center", wrap_text=False)
    if borde:
        lado = Side(style="thin", color=COLOR_BORDE)
        from openpyxl.styles import Border as Brd
        cell.border = Brd(left=lado, right=lado, top=lado, bottom=lado)
    return cell


def _hoja_resumen(wb, consejos: dict, fecha: str):
    """Crea la hoja de portada con estadísticas generales."""
    from openpyxl.styles import Font, PatternFill, Alignment

    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False

    # Anchos de columna
    for col, ancho in [("A",4),("B",40),("C",16),("D",16),("E",16),("F",16)]:
        ws.column_dimensions[col].width = ancho

    # Título
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 32
    ws.merge_cells("B2:F2")
    c = ws["B2"]
    c.value = "PADRÓN DE VOCEROS — CONSEJOS COMUNALES"
    c.font = Font(name=FUENTE, bold=True, size=16, color=COLOR_BLANCO)
    c.fill = PatternFill("solid", fgColor=COLOR_HEADER)
    c.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[3].height = 20
    ws.merge_cells("B3:F3")
    c = ws["B3"]
    c.value = "Comuna Productiva Presidente Obrero"
    c.font = Font(name=FUENTE, size=11, color=COLOR_BLANCO)
    c.fill = PatternFill("solid", fgColor=COLOR_SUBHEADER)
    c.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[4].height = 16
    ws.merge_cells("B4:F4")
    c = ws["B4"]
    c.value = f"Generado el {fecha}"
    c.font = Font(name=FUENTE, size=9, color="666666")
    c.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[5].height = 12

    # KPIs
    total_voceros     = sum(len(d["voceros"]) for d in consejos.values())
    total_consejos    = len(consejos)
    total_principales = sum(
        sum(1 for v in d["voceros"] if v["tipo"] == "Principal")
        for d in consejos.values()
    )
    total_suplentes = total_voceros - total_principales

    ws.row_dimensions[6].height = 20
    ws.row_dimensions[7].height = 36
    for col, (label, valor, color) in enumerate([
        ("Total Consejos",    total_consejos,    COLOR_HEADER),
        ("Total Voceros",     total_voceros,     COLOR_SUBHEADER),
        ("Principales",       total_principales, "1A7A45"),
        ("Suplentes",         total_suplentes,   "8E44AD"),
    ], start=2):
        c = ws.cell(row=6, column=col, value=label)
        c.font = Font(name=FUENTE, bold=True, size=9, color=COLOR_BLANCO)
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c2 = ws.cell(row=7, column=col, value=valor)
        c2.font = Font(name=FUENTE, bold=True, size=20, color=color)
        c2.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[8].height = 12

    # Tabla de consejos
    headers = ["#", "Consejo Comunal", "Sector", "Voceros", "Principales", "Suplentes"]
    ws.row_dimensions[9].height = 20
    for col, h in enumerate(headers, start=2):
        _celda(ws, 9, col, h, negrita=True, color_fondo=COLOR_HEADER,
               color_fuente=COLOR_BLANCO, tamaño=10, alineacion="center", borde=True)

    for i, (nombre_c, datos_c) in enumerate(sorted(consejos.items()), start=1):
        fila = 9 + i
        ws.row_dimensions[fila].height = 16
        vocs = datos_c["voceros"]
        princ = sum(1 for v in vocs if v["tipo"] == "Principal")
        supl  = len(vocs) - princ
        bg = COLOR_ALT_ROW if i % 2 == 0 else COLOR_BLANCO
        for col, (val, alin) in enumerate(zip(
            [i, nombre_c, datos_c["sector"], len(vocs), princ, supl],
            ["center","left","left","center","center","center"]
        ), start=2):
            _celda(ws, fila, col, val, color_fondo=bg, tamaño=10, alineacion=alin, borde=True)

    # Nota final
    fila_nota = 10 + len(consejos)
    ws.row_dimensions[fila_nota].height = 14
    c = ws.cell(row=fila_nota, column=2,
                value="→ Ver pestañas individuales por Consejo Comunal")
    c.font = Font(name=FUENTE, italic=True, size=9, color="888888")


def _hoja_consejo(wb, nombre_consejo: str, datos: dict, numero: int):
    """Crea una hoja individual por consejo comunal."""
    from openpyxl.styles import Font, PatternFill, Alignment

    nombre_hoja = re.sub(r'[:/\\?*\[\]]', '', nombre_consejo)[:28].strip() or f"Consejo_{numero}"
    ws = wb.create_sheet(title=nombre_hoja)
    ws.sheet_view.showGridLines = False

    for col, ancho in [("A",4),("B",4),("C",30),("D",14),("E",28),("F",14),("G",14),("H",8)]:
        ws.column_dimensions[col].width = ancho

    # Encabezado
    ws.merge_cells("B1:H1")
    c = ws["B1"]
    c.value = nombre_consejo.upper()
    c.font = Font(name=FUENTE, bold=True, size=13, color=COLOR_BLANCO)
    c.fill = PatternFill("solid", fgColor=COLOR_HEADER)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells("B2:H2")
    c = ws["B2"]
    sector = datos.get("sector", "")
    c.value = f"Sector: {sector}" if sector else "Sector: —"
    c.font = Font(name=FUENTE, size=9, color="555555")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 14
    ws.row_dimensions[3].height = 8

    # Cabeceras de tabla
    columnas = ["#", "Nombre Completo", "C.I.", "Teléfono", "Vocería / Cargo", "Tipo", "Votos"]
    ws.row_dimensions[4].height = 18
    for col, label in enumerate(columnas, start=2):
        _celda(ws, 4, col, label, negrita=True, color_fondo=COLOR_SUBHEADER,
               color_fuente=COLOR_BLANCO, tamaño=10, alineacion="center", borde=True)

    ws.auto_filter.ref = "B4:H4"

    # Filas de voceros
    voceros = datos["voceros"]
    for i, v in enumerate(voceros, start=1):
        fila = 4 + i
        ws.row_dimensions[fila].height = 15
        es_suplente = v["tipo"].lower() == "suplente"
        bg = "FFF8F0" if es_suplente else (COLOR_ALT_ROW if i % 2 == 0 else COLOR_BLANCO)

        valores = [i, v["nombre"], v["cedula"] or "—", v["telefono"] or "—",
                   v["cargo"], v["tipo"], v["votos"] if v["votos"] else "—"]
        alineaciones = ["center","left","center","center","left","center","center"]

        for col, (val, alin) in enumerate(zip(valores, alineaciones), start=2):
            cell = _celda(ws, fila, col, val, color_fondo=bg, tamaño=10, alineacion=alin, borde=True)
            if es_suplente:
                cell.font = Font(name=FUENTE, size=10, italic=True, color="555555")

    # Fila de totales
    fila_tot = 4 + len(voceros) + 1
    ws.row_dimensions[fila_tot].height = 16
    ws.merge_cells(f"B{fila_tot}:D{fila_tot}")
    c = ws.cell(row=fila_tot, column=2, value=f"Total: {len(voceros)} voceros")
    c.font = Font(name=FUENTE, bold=True, size=10, color=COLOR_BLANCO)
    c.fill = PatternFill("solid", fgColor=COLOR_ACENTO)
    c.alignment = Alignment(horizontal="center", vertical="center")

    princ = sum(1 for v in voceros if v["tipo"] == "Principal")
    supl  = len(voceros) - princ
    ws.cell(row=fila_tot, column=5, value=f"Principales: {princ}").font = Font(name=FUENTE, bold=True, size=9)
    ws.cell(row=fila_tot, column=6, value=f"Suplentes: {supl}").font   = Font(name=FUENTE, bold=True, size=9)


async def exportar_padron_excel(output_path: Optional[Path] = None) -> bytes:
    """
    Genera el Excel del padrón completo.
    Retorna bytes del archivo .xlsx para enviarlo desde FastAPI como StreamingResponse.
    """
    import openpyxl

    fecha_gen     = datetime.now().strftime("%d/%m/%Y %H:%M")
    consejos      = await _obtener_datos_padron()

    if not consejos:
        raise ValueError("No hay voceros registrados en el padrón para exportar.")

    wb = openpyxl.Workbook()
    _hoja_resumen(wb, consejos, fecha_gen)

    for numero, (nombre_c, datos_c) in enumerate(sorted(consejos.items()), start=1):
        _hoja_consejo(wb, nombre_c, datos_c, numero)

    wb.properties.title       = "Padrón de Voceros"
    wb.properties.subject     = "Comuna Productiva Presidente Obrero"
    wb.properties.creator     = "Asistente Comunal"
    wb.properties.description = f"Generado el {fecha_gen}"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    xlsx_bytes = buffer.read()

    if output_path:
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        Path(output_path).write_bytes(xlsx_bytes)
        logger.info(f"Padrón exportado a: {output_path}")

    total = sum(len(d["voceros"]) for d in consejos.values())
    logger.info(f"Excel generado: {len(consejos)} consejos, {total} voceros, {len(xlsx_bytes)//1024} KB")
    return xlsx_bytes